"""Development-only extraction of the supplied NBC workbook; no application importer.

Usage: python scripts/nbc-bp03-extract.py SOURCE.xlsx [OUTPUT_DIRECTORY]
Requires openpyxl. Outputs contain private BOQ data and must stay outside Git.
"""
import collections
import hashlib
import json
import math
from pathlib import Path
import re
import sys

import openpyxl


def extract(source):
    formulas = openpyxl.load_workbook(source, data_only=False)
    values = openpyxl.load_workbook(source, data_only=True)
    sheet = values['Progress']
    assert values.sheetnames == ['Progress', 'Sitha', 'Lychhinh', 'Chhorvaon', 'Srun', 'Kheng', 'Tola']
    errors = [(s.title, c.coordinate, c.value) for s in values for row in s for c in row if c.data_type == 'e']
    assert not errors, errors
    assigned = {}
    for handler in values.sheetnames[1:]:
        for row in values[handler].iter_rows(min_row=5):
            wbs, previous, actual = row[0].value, row[10].value, row[12].value
            if wbs and actual is not None:
                assert wbs not in assigned, f'Duplicate assignment: {wbs}'
                assert isinstance(previous, (int, float)) and isinstance(actual, (int, float))
                assigned[wbs] = (handler, previous, actual)
    nodes = []
    for r in range(6, sheet.max_row + 1):
        wbs = sheet.cell(r, 1).value
        if not wbs:
            continue
        assert wbs == 'GR&P' or re.fullmatch(r'1\.6(?:\.\d+)+', wbs), wbs
        parent = wbs.rsplit('.', 1)[0] if wbs.count('.') > 2 else None
        hierarchy = [sheet.cell(r, c).value for c in range(2, 10)]
        nodes.append(dict(
            id='nbc:' + wbs, projectId='nbc-bp03', wbs=wbs,
            parentId='nbc:' + parent if parent else None, order=r,
            name=sheet.cell(r, 10).value, description=' / '.join(str(v) for v in hierarchy if v),
            relativeWeight=sheet.cell(r, 11).value, globalWeight=sheet.cell(r, 19).value,
            baseline=sheet.cell(r, 14).value, progress=sheet.cell(r, 14).value,
            sourcePrevious=sheet.cell(r, 12).value, sourceWeekly=sheet.cell(r, 13).value,
            duration=sheet.cell(r, 18).value, boqValue=sheet.cell(r, 20).value,
            schedule=sheet.cell(r, 21).value, handler=assigned.get(wbs, ('',))[0],
            remark='', history=[], sourceRow=r,
            sourceFormulas={sheet.cell(4,c).value or str(c): formulas['Progress'].cell(r,c).value
                            for c in range(11,22) if formulas['Progress'].cell(r,c).data_type == 'f'},
        ))
    by_id = {n['id']: n for n in nodes}
    assert len(by_id) == len(nodes) == 790
    children = collections.defaultdict(list)
    for n in nodes:
        assert n['parentId'] is None or n['parentId'] in by_id
        children[n['parentId']].append(n)
    roots = children[None]
    total_boq = sum(n['boqValue'] or 0 for n in roots)
    checks = []
    def close(actual, expected, label, tolerance=1e-11):
        assert isinstance(expected, (int, float)) and math.isfinite(expected), (label, expected)
        assert abs(actual - expected) <= tolerance, (label, actual, expected)
    for n in nodes:
        cs = children[n['id']]
        n['editable'] = not cs
        root = n
        lineage = []
        while root['parentId']:
            lineage.insert(0, root)
            root = by_id[root['parentId']]
        n['area'] = root['name']
        n['location'] = root['name'].split(' - ')[0] if ' - ' in root['name'] else root['name']
        n['trade'] = lineage[0]['name'] if lineage else root['name']
        if cs:
            duration = sum(c['duration'] or 0 for c in cs)
            close(duration, n['duration'], n['wbs'] + ' duration')
            weight_sum = sum(c['relativeWeight'] for c in cs)
            progress = sum(c['relativeWeight'] * c['progress'] for c in cs) / weight_sum if weight_sum else 0
            close(progress, n['progress'], n['wbs'] + ' parent')
            expected_refs = {c['sourceRow'] for c in cs}
            refs = {int(v) for v in re.findall(r'K(\d+)\*N\d+', formulas['Progress'].cell(n['sourceRow'],14).value)}
            assert refs == expected_refs, (n['wbs'], refs, expected_refs)
            checks.append(n['wbs'])
        else:
            assert 0 <= n['progress'] <= 1
            if n['wbs'] != 'GR&P':
                handler, previous, actual = assigned[n['wbs']]
                close(actual, n['progress'], n['wbs'] + ' handler actual')
                close(previous, n['sourcePrevious'], n['wbs'] + ' handler previous')
                assert handler + '!' in formulas['Progress'].cell(n['sourceRow'],14).value
        if n['parentId']:
            siblings = children[n['parentId']]
            duration = sum(c['duration'] or 0 for c in siblings)
            close((n['duration'] or 0) / duration if duration else 0, n['relativeWeight'], n['wbs'] + ' relative')
            close(n['relativeWeight'] * by_id[n['parentId']]['globalWeight'], n['globalWeight'], n['wbs'] + ' global')
        else:
            close(n['boqValue'] / total_boq, n['relativeWeight'], n['wbs'] + ' root')
            close(n['relativeWeight'], n['globalWeight'], n['wbs'] + ' root global')
    counts = dict(collections.Counter(n['handler'] for n in nodes if n['handler']))
    assert counts == dict(Sitha=193, Lychhinh=50, Chhorvaon=6, Srun=31, Kheng=31, Tola=250), counts
    leaves = [n for n in nodes if n['editable']]
    assert len(leaves) == len(assigned) + 1 == 562
    weight = sum(n['globalWeight'] for n in leaves)
    overall = sum(n['globalWeight'] * n['progress'] for n in leaves)
    close(weight, sheet['K5'].value, 'total weight')
    close(total_boq, sheet['T5'].value, 'total BOQ', 1e-7)
    close(overall, sheet['N5'].value, 'overall')
    # Independent propagation check at every ancestor of a changed assigned leaf.
    leaf = next(n for n in leaves if n['handler'])
    changed = {n['id']: n['progress'] for n in nodes}
    changed[leaf['id']] = 0.375
    for n in reversed(nodes):
        cs = children[n['id']]
        if cs:
            denominator = sum(c['relativeWeight'] for c in cs)
            changed[n['id']] = sum(c['relativeWeight'] * changed[c['id']] for c in cs) / denominator if denominator else 0
    propagated = sum(n['globalWeight'] * changed[n['id']] for n in roots)
    close(propagated, overall + leaf['globalWeight'] * (0.375 - leaf['progress']), 'edited propagation')
    project = dict(id='nbc-bp03', name='NBC BP03 – Pan Pacific Hotel', workspace='team',
                   color='blue', description='Interior Fit-Out · NBC BP03', target='', status='Active',
                   progress=overall, nbcBp03=True)
    reconciliation = dict(nodes=len(nodes), assignedActivities=len(assigned), editableEntries=len(leaves),
                          handlers=counts, rootWeight=sum(n['relativeWeight'] for n in roots),
                          effectiveWeight=weight, boqValue=total_boq, overall=overall,
                          workbookOverall=sheet['N5'].value, parentsVerified=len(checks),
                          areas=[dict(wbs=n['wbs'], name=n['name'], progress=n['progress']) for n in roots],
                          sourceSha256=hashlib.sha256(Path(source).read_bytes()).hexdigest(), reportingDate=None)
    return dict(project=project, activities=nodes, reconciliation=reconciliation)


def seed_sql(dataset):
    # psql supplies owner_id explicitly; this file is private, reviewed, and never run by app startup.
    payload = json.dumps(dataset, ensure_ascii=False, separators=(',', ':')).replace("'", "''")
    return """\\set ON_ERROR_STOP on
begin;
select set_config('nbc.owner_id', :'owner_id', true);
select set_config('request.jwt.claim.sub', :'owner_id', true);
do $nbc_seed$
declare owner uuid := current_setting('nbc.owner_id')::uuid; payload jsonb := '""" + payload + """'::jsonb;
pid text; matches integer; total integer;
begin
  if not exists(select 1 from auth.users where id=owner) then raise exception 'Owner account does not exist'; end if;
  perform pg_advisory_xact_lock(hashtextextended('nbc-bp03-initialization',0));
  if exists(select 1 from public.kw_projects where data->>'nbcBp03'='true' and user_id<>owner) then raise exception 'Company NBC is already initialized under another storage owner; use that existing owner instead of creating a copy'; end if;
  select count(*),min(id) into matches,pid from public.kw_projects
    where user_id=owner and (data->>'nbcBp03'='true' or regexp_replace(lower(data->>'name'),'[^a-z0-9]','','g') like 'nbcbp03%');
  if matches>1 then raise exception 'Multiple NBC BP03 projects; resolve explicitly before initializing'; end if;
  if matches=0 then
    pid='nbc-bp03';
    if exists(select 1 from public.kw_projects where user_id=owner and id=pid) then raise exception 'Project ID is already used'; end if;
    insert into public.kw_projects(user_id,id,data) values(owner,pid,payload->'project');
  else
    update public.kw_projects set data=data||jsonb_build_object('nbcBp03',true,'workspace','team'),updated_at=clock_timestamp()
    where user_id=owner and id=pid and (data->>'nbcBp03' is distinct from 'true' or data->>'workspace' is distinct from 'team');
  end if;
  insert into public.kw_nbc_activities(user_id,id,data)
    select owner,a->>'id',a||jsonb_build_object('projectId',pid) from jsonb_array_elements(payload->'activities') a
    on conflict(user_id,id) do nothing;
  select count(*) into total from public.kw_nbc_activities where user_id=owner and project_id=pid;
  if total<>jsonb_array_length(payload->'activities') then raise exception 'Unexpected WBS count; initialization rolled back'; end if;
  if exists(select 1 from public.kw_nbc_activities a where a.user_id=owner and a.project_id=pid and
    ((a.data->>'editable')::boolean is distinct from not exists(select 1 from public.kw_nbc_activities c where c.user_id=owner and c.parent_id=a.id))) then
    raise exception 'Invalid WBS editable rule';
  end if;
  -- Existing progress, assignments, remarks, history and reports are never updated.
end $nbc_seed$;
commit;
"""


if __name__ == '__main__':
    output = Path(sys.argv[2] if len(sys.argv) > 2 else '.private/nbc-bp03')
    output.mkdir(parents=True, exist_ok=True)
    dataset = extract(sys.argv[1])
    (output / 'initial.json').write_text(json.dumps(dataset, ensure_ascii=False, indent=2), encoding='utf-8')
    (output / 'reconciliation.json').write_text(json.dumps(dataset['reconciliation'], ensure_ascii=False, indent=2), encoding='utf-8')
    (output / 'initialize.sql').write_text(seed_sql(dataset), encoding='utf-8')
    print(json.dumps({k:v for k,v in dataset['reconciliation'].items() if k != 'areas'}, indent=2))
